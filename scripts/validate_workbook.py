#!/usr/bin/env python3
"""Read-only structural and freeze-gate checks for the v4.2 LCI template."""

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_SHEETS = {
    "00_工作流总控", "项目说明", "数据来源", "候选数据", "模型过程",
    "模型参数", "模型流与公式", "核算检查", "问题与决策", "绘图确认",
    "参数覆盖度矩阵", "待补参数清单",
}

REQUIRED_HEADERS = {
    "数据来源": {"来源ID", "来源体系ID", "文件角色", "原研究方法", "文献池角色", "完整清单证据"},
    "候选数据": {"候选ID", "目标对象类型", "目标对象ID", "来源ID", "原始值", "原始单位", "原始基准", "可比性", "采用状态"},
    "模型参数": {"参数ID", "采用候选ID"},
    "模型流与公式": {"流ID", "采用候选ID", "采用计量基准", "方法采用状态"},
    "核算检查": {"检查ID", "审查路线", "实际验证结论", "影响对象ID"},
    "参数覆盖度矩阵": {"覆盖ID", "模块/过程ID", "操作ID", "所需参数/流ID", "数据层级", "覆盖状态", "绘图可用层级"},
    "待补参数清单": {"待补ID", "归属操作ID", "参数/流ID", "关联问题ID", "用户决定", "状态"},
}


def value(cell):
    raw = cell.value
    return raw.strip() if isinstance(raw, str) else raw


def headers(sheet):
    return {value(cell): cell.column for cell in sheet[4] if value(cell)}


def rows(sheet):
    h = headers(sheet)
    for row in sheet.iter_rows(min_row=5, max_col=max(h.values(), default=1)):
        if value(row[0]) in (None, ""):
            continue
        yield row[0].row, {name: value(row[index - 1]) for name, index in h.items()}


def split_ids(text):
    if not text:
        return []
    return [part.strip() for part in str(text).replace(",", ";").split(";") if part.strip()]


def has_complete_lci_proof(system_rows):
    """An eligible framework has a pool role and an explicit whole-chain inventory locator."""
    return any(
        "完整清单文献池" in split_ids(record.get("文献池角色"))
        and record.get("完整清单证据")
        for _, record in system_rows
    )


def unresolved_issue_needs_decision(record):
    return record.get("状态") not in {"已解决", "resolved", "closed"} and not record.get("用户决定")


def validate(path: Path, freeze: bool):
    errors = []
    warnings = []
    wb = load_workbook(path, read_only=True, data_only=False)
    missing_sheets = sorted(REQUIRED_SHEETS - set(wb.sheetnames))
    if missing_sheets:
        return [f"Missing sheets: {', '.join(missing_sheets)}"], warnings

    for sheet_name, expected in REQUIRED_HEADERS.items():
        missing = sorted(expected - set(headers(wb[sheet_name])))
        if missing:
            errors.append(f"{sheet_name}: missing headers: {', '.join(missing)}")
    if errors:
        return errors, warnings

    source_rows = list(rows(wb["数据来源"]))
    source_ids = [record.get("来源ID") for _, record in source_rows]
    if len(source_ids) != len(set(source_ids)):
        errors.append("数据来源: duplicate 来源ID")
    source_systems = {}
    for row_num, record in source_rows:
        system_id = record.get("来源体系ID")
        if system_id:
            source_systems.setdefault(system_id, []).append((row_num, record))
        if freeze and record.get("文献池角色"):
            roles = split_ids(record.get("文献池角色"))
            invalid = set(roles) - {"完整清单文献池", "参数选取文献池"}
            if invalid:
                errors.append(f"数据来源 row {row_num}: unknown 文献池角色 {sorted(invalid)}")

    candidate_rows = list(rows(wb["候选数据"]))
    candidate_ids = [record.get("候选ID") for _, record in candidate_rows]
    if len(candidate_ids) != len(set(candidate_ids)):
        errors.append("候选数据: duplicate 候选ID")
    candidates = {record.get("候选ID"): record for _, record in candidate_rows}
    for row_num, record in candidate_rows:
        if record.get("来源ID") not in source_ids:
            errors.append(f"候选数据 row {row_num}: 来源ID not found in 数据来源")
        if record.get("目标对象类型") not in {"参数", "流"}:
            errors.append(f"候选数据 row {row_num}: 目标对象类型 must be 参数 or 流")
        if not record.get("证据位置"):
            errors.append(f"候选数据 row {row_num}: missing 证据位置")
        if record.get("原始值") in (None, "") or not record.get("原始单位") or not record.get("原始基准"):
            errors.append(f"候选数据 row {row_num}: original value, unit and basis are required")
        if record.get("换算值") not in (None, "") and not record.get("换算依据"):
            errors.append(f"候选数据 row {row_num}: converted value requires 换算依据")
        if record.get("可比性") == "不可直接比较" and record.get("采用状态") == "采用":
            errors.append(f"候选数据 row {row_num}: non-comparable candidate cannot be silently adopted")

    parameters = list(rows(wb["模型参数"]))
    flows = list(rows(wb["模型流与公式"]))
    processes = list(rows(wb["模型过程"]))
    process_ids = [record.get("过程ID") for _, record in processes]
    if len(process_ids) != len(set(process_ids)):
        errors.append("模型过程: duplicate 过程ID")
    if len([record.get("流ID") for _, record in flows]) != len({record.get("流ID") for _, record in flows}):
        errors.append("模型流与公式: duplicate 流ID")
    parameter_ids = {record.get("参数ID") for _, record in parameters}
    flow_ids = {record.get("流ID") for _, record in flows}
    coverage_rows = list(rows(wb["参数覆盖度矩阵"]))
    coverage_ids = [record.get("覆盖ID") for _, record in coverage_rows]
    if len(coverage_ids) != len(set(coverage_ids)):
        errors.append("参数覆盖度矩阵: duplicate 覆盖ID")
    known_operation_ids = set()
    for row_num, record in coverage_rows:
        if record.get("模块/过程ID") not in process_ids:
            errors.append(f"参数覆盖度矩阵 row {row_num}: unknown 模块/过程ID")
        if not record.get("操作ID"):
            errors.append(f"参数覆盖度矩阵 row {row_num}: 操作ID is required")
        else:
            known_operation_ids.add(record.get("操作ID"))
        if record.get("数据层级") not in {"操作级直接给出", "公式唯一拆出", "过程/模块总量", "仅拓扑", "缺失"}:
            errors.append(f"参数覆盖度矩阵 row {row_num}: invalid 数据层级")
        if not record.get("覆盖状态") or not record.get("绘图可用层级"):
            errors.append(f"参数覆盖度矩阵 row {row_num}: coverage and drawing levels required")
    issue_rows = list(rows(wb["问题与决策"]))
    issue_ids = {record.get("问题ID") for _, record in issue_rows}
    pending_rows = list(rows(wb["待补参数清单"]))
    pending_ids = [record.get("待补ID") for _, record in pending_rows]
    if len(pending_ids) != len(set(pending_ids)):
        errors.append("待补参数清单: duplicate 待补ID")
    for row_num, record in pending_rows:
        if record.get("归属操作ID") not in known_operation_ids:
            errors.append(f"待补参数清单 row {row_num}: 归属操作ID not in 参数覆盖度矩阵")
        if record.get("关联问题ID") not in issue_ids:
            errors.append(f"待补参数清单 row {row_num}: 关联问题ID not found in 问题与决策")
        if freeze and (not record.get("用户决定") or not record.get("状态")):
            errors.append(f"待补参数清单 row {row_num}: user decision and status required before freeze")
    used_candidates = set()
    for sheet_name, model_rows, id_name in (
        ("模型参数", parameters, "参数ID"),
        ("模型流与公式", flows, "流ID"),
    ):
        for row_num, record in model_rows:
            ids = split_ids(record.get("采用候选ID"))
            for candidate_id in ids:
                used_candidates.add(candidate_id)
                candidate = candidates.get(candidate_id)
                if candidate is None:
                    errors.append(f"{sheet_name} row {row_num}: unknown candidate {candidate_id}")
                    continue
                if candidate.get("目标对象ID") != record.get(id_name):
                    errors.append(f"{sheet_name} row {row_num}: candidate {candidate_id} points to another model object")
                expected_kind = "参数" if sheet_name == "模型参数" else "流"
                if candidate.get("目标对象类型") != expected_kind:
                    errors.append(f"{sheet_name} row {row_num}: candidate {candidate_id} has wrong target type")
                if candidate.get("采用状态") != "采用":
                    errors.append(f"{sheet_name} row {row_num}: candidate {candidate_id} is not marked 采用")
            if freeze and not ids:
                if sheet_name == "模型参数":
                    errors.append(f"模型参数 row {row_num}: adopted parameter lacks 采用候选ID")
                elif not record.get("采用公式") or not record.get("输入参数ID"):
                    errors.append(f"模型流与公式 row {row_num}: flow needs adopted candidate or formula plus input parameters")
            if freeze and sheet_name == "模型流与公式":
                if not record.get("方向") or not record.get("采用单位") or not record.get("采用计量基准"):
                    errors.append(f"模型流与公式 row {row_num}: direction, adopted unit and basis are required")

    for candidate_id, record in candidates.items():
        if record.get("采用状态") == "采用" and candidate_id not in used_candidates:
            errors.append(f"候选数据 {candidate_id}: marked 采用 but no model record references it")
        kind = record.get("目标对象类型")
        target = record.get("目标对象ID")
        if kind == "参数" and target not in parameter_ids:
            warnings.append(f"候选数据 {candidate_id}: target parameter not yet in adopted model")
        if kind == "流" and target not in flow_ids:
            warnings.append(f"候选数据 {candidate_id}: target flow not yet in adopted model")

    project = wb["项目说明"]
    fields = {value(project.cell(row, 1)): value(project.cell(row, 2)) for row in range(4, 23)}
    if freeze:
        for label in ("目标产品系统", "功能单位", "系统边界", "产品质量规格"):
            if not fields.get(label):
                errors.append(f"项目说明: missing {label}")
        if fields.get("建模模式") not in {"attributional", "consequential"}:
            errors.append("项目说明: 建模模式 must be attributional or consequential")
        if fields.get("工作簿版本") != "4.2":
            errors.append("项目说明: 工作簿版本 must be 4.2 for this workflow")
        if fields.get("工艺路线") != "主文献锚定":
            errors.append("项目说明: 工艺路线 must be 主文献锚定; no automatic pure-model fallback")
        if not fields.get("整链主来源体系ID"):
            errors.append("项目说明: 主文献锚定 needs 整链主来源体系ID")
        if fields.get("整链主来源体系ID") not in source_systems:
            errors.append("项目说明: 整链主来源体系ID not found in 数据来源")
        else:
            selected = source_systems[fields.get("整链主来源体系ID")]
            if not has_complete_lci_proof(selected):
                errors.append("项目说明: framework source must have complete-LCI pool role and evidence")
        if not coverage_rows:
            errors.append("参数覆盖度矩阵: at least one operation-level coverage row is required")
        missing_processes = set(process_ids) - {record.get("模块/过程ID") for _, record in coverage_rows}
        if missing_processes:
            errors.append(f"参数覆盖度矩阵: processes without coverage rows: {sorted(missing_processes)}")
        for row_num, record in processes:
            if not record.get("工艺骨架依据") or not record.get("骨架确认决策ID"):
                errors.append(f"模型过程 row {row_num}: evidence and user-confirmed skeleton decision required")
        if not processes or not flows:
            errors.append("A frozen foreground model needs at least one process and one flow")
        value_wb = load_workbook(path, read_only=True, data_only=True)
        value_flows = dict(rows(value_wb["模型流与公式"]))
        for row_num, record in flows:
            kind = record.get("值类型")
            if kind not in {"physical_material", "physical_energy", "environment_exchange", "allocated_burden", "substitution_candidate"}:
                errors.append(f"模型流与公式 row {row_num}: 值类型 unknown or missing")
            if kind in {"physical_material", "physical_energy", "environment_exchange"}:
                result = value_flows.get(row_num, {}).get("图示代表值")
                if not isinstance(result, (int, float)) or isinstance(result, bool):
                    errors.append(f"模型流与公式 row {row_num}: 图示代表值 must be a recalculated number")
                source_id = record.get("来源过程ID")
                target_id = record.get("目标过程ID")
                owner = record.get("过程ID")
                if not source_id and record.get("方向") == "输入" and target_id == owner:
                    source_id = "EXT_IN"
                if not target_id and record.get("方向") == "输出" and source_id == owner:
                    target_id = "EXT_OUT"
                if source_id not in set(process_ids) | {"EXT_IN"} or target_id not in set(process_ids) | {"EXT_OUT"}:
                    errors.append(f"模型流与公式 row {row_num}: source/target process is not resolved")
        control = wb["00_工作流总控"]
        for row_num in range(5, 14):
            gate = value(control.cell(row_num, 1))
            if value(control.cell(row_num, 6)) not in {"完成", "已确认"}:
                errors.append(f"00_工作流总控 {gate}: stage must be 完成 or 已确认 before freeze")
            if row_num in {5, 8, 13} and not value(control.cell(row_num, 7)):
                errors.append(f"00_工作流总控 {gate}: confirmation person/record missing")
        for label in ("来源路线确认记录", "数据功能用户确认", "冻结版本"):
            if not fields.get(label):
                errors.append(f"项目说明: missing {label}")
        if fields.get("数据功能状态") != "frozen":
            errors.append("项目说明: 数据功能状态 is not frozen")
        checks = list(rows(wb["核算检查"]))
        if not checks:
            errors.append("核算检查: at least one executed or explicitly not-calculated check is required")
        for row_num, record in checks:
            if record.get("结果") == "BLOCKED":
                errors.append(f"核算检查 row {row_num}: unresolved BLOCKED check")
            if not record.get("结果") or not record.get("实际验证结论"):
                errors.append(f"核算检查 row {row_num}: result and actual conclusion are required")
        for row_num, record in issue_rows:
            if record.get("是否阻断冻结") in {"是", "YES", True} and record.get("状态") not in {"已解决", "resolved", "closed"}:
                errors.append(f"问题与决策 row {row_num}: blocking issue unresolved")
            if unresolved_issue_needs_decision(record):
                errors.append(f"问题与决策 row {row_num}: unresolved post-QC issue needs a user decision before freeze")
        if not parameters and not flows:
            errors.append("No adopted parameters or flows are recorded")
    else:
        warnings.append("Draft/structure check only; use --freeze after the user has reviewed the completed workbook")

    return errors, warnings


def main():
    parser = argparse.ArgumentParser(description="Validate a v4.2 LCI workbook without modifying it")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--freeze", action="store_true", help="Check the user-confirmed freeze gate")
    args = parser.parse_args()
    try:
        errors, warnings = validate(args.workbook, args.freeze)
    except FileNotFoundError:
        print(f"ERROR: workbook not found: {args.workbook}", file=sys.stderr)
        return 1
    for message in warnings:
        print(f"WARNING: {message}")
    if errors:
        for message in errors:
            print(f"ERROR: {message}", file=sys.stderr)
        print(f"FAILED: {len(errors)} error(s)", file=sys.stderr)
        return 1
    print(f"PASS: workbook {'freeze gate' if args.freeze else 'structure'} checks: {args.workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
