#!/usr/bin/env python3
"""Compile a draft drawing spec from a frozen workbook and audit scientific fields.

This is a read-only workbook consumer. It never calculates a new LCI value.
"""

import argparse
import hashlib
import json
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook

from validate_workbook import validate as validate_data


PHYSICAL = {"physical_material", "physical_energy", "environment_exchange"}
NONPHYSICAL = {"allocated_burden", "substitution_candidate"}


def cell_value(cell):
    value = cell.value
    return value.strip() if isinstance(value, str) else value


def records(sheet):
    keys = {cell.column: cell_value(cell) for cell in sheet[4] if cell_value(cell)}
    for row in sheet.iter_rows(min_row=5, max_col=max(keys, default=1)):
        if cell_value(row[0]) not in (None, ""):
            yield row[0].row, {name: cell_value(row[col - 1]) for col, name in keys.items()}


def project_fields(sheet):
    return {cell_value(sheet.cell(row, 1)): cell_value(sheet.cell(row, 2)) for row in range(4, 23)}


FINGERPRINT_SHEETS = (
    "00_工作流总控", "项目说明", "数据来源", "候选数据", "模型过程", "模型参数",
    "模型流与公式", "核算检查", "问题与决策", "绘图确认",
    "参数覆盖度矩阵", "待补参数清单",
)


def content_hash(values_book, formulas_book):
    """Hash the actual fields and calculated caches, not ZIP timestamps/styles."""
    digest = hashlib.sha256()
    for name in FINGERPRINT_SHEETS:
        values = values_book[name]
        formulas = formulas_book[name]
        digest.update(name.encode("utf-8"))
        for formula_row, value_row in zip(formulas.iter_rows(), values.iter_rows()):
            cells = []
            for formula_cell, value_cell in zip(formula_row, value_row):
                cells.append((cell_value(formula_cell), cell_value(value_cell))
                             if formula_cell.data_type == "f" else cell_value(formula_cell))
            digest.update(json.dumps(cells, ensure_ascii=False, default=str, separators=(",", ":")).encode("utf-8"))
    return digest.hexdigest()


def load_snapshot(path):
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    wf = load_workbook(path, read_only=True, data_only=False)
    drawing_rows = list(records(wb["绘图确认"]))
    drawing_formula_rows = list(records(wf["绘图确认"]))
    drawing_ids = [row["对象ID"] for _, row in drawing_rows]
    if len(drawing_ids) != len(set(drawing_ids)):
        raise ValueError("绘图确认存在重复对象ID；不能按最后一行覆盖前面的决定")
    snapshot = {
        "project": project_fields(wb["项目说明"]),
        "processes": dict((row["过程ID"], row) for _, row in records(wb["模型过程"])),
        "flows": dict((row["流ID"], row) for _, row in records(wb["模型流与公式"])),
        "drawing": dict((row["对象ID"], row) for _, row in drawing_rows),
        "drawing_formulas": dict((row["对象ID"], row) for _, row in drawing_formula_rows),
        "sha256": content_hash(wb, wf),
    }
    return snapshot


def endpoints(flow, process_ids):
    source = flow.get("来源过程ID")
    target = flow.get("目标过程ID")
    owner = flow.get("过程ID")
    direction = flow.get("方向")
    if not source and direction == "输入" and target == owner:
        source = "EXT_IN"
    if not target and direction == "输出" and source == owner:
        target = "EXT_OUT"
    if not source or not target:
        raise ValueError(f"流 {flow.get('流ID')}: 起止过程不明确；先在 Excel 确认端点")
    for endpoint in (source, target):
        if endpoint not in process_ids | {"EXT_IN", "EXT_OUT"}:
            raise ValueError(f"流 {flow.get('流ID')}: 端点 {endpoint} 不在模型过程表")
    return source, target


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def compile_spec(workbook, template):
    errors, _ = validate_data(Path(workbook), freeze=True)
    if errors:
        raise ValueError("数据功能尚未通过冻结检查：" + "；".join(errors[:5]))
    source = load_snapshot(workbook)
    spec = deepcopy(template)
    spec["schema_version"] = "0.8"
    p = source["project"]
    spec["spec_id"] = "draft-" + source["sha256"][:12]
    spec["metadata"].update({
        "product_system": p.get("目标产品系统"),
        "product_name": {"zh": p.get("产品名称"), "en": None},
        "functional_unit": p.get("功能单位"),
        "system_boundary": p.get("系统边界"),
        "modeling_mode": p.get("建模模式"),
        "inventory_version": p.get("冻结版本"),
        "source_workbook": str(Path(workbook).resolve()),
        "source_content_sha256": source["sha256"],
    })
    spec["confirmations"]["data"] = {"status": "confirmed", "note": p.get("数据功能用户确认")}
    spec["confirmations"]["figure"] = {"status": "draft", "note": None}
    process_ids = set(source["processes"])
    if not process_ids:
        raise ValueError("模型过程为空，不能编译绘图蓝图")
    spec["nodes"] = [
        {
            "node_id": process_id,
            "process_id": process_id,
            "workshop_id": row.get("功能车间"),
            "name": row.get("中文名称") or row.get("模型过程原名"),
            "sequence": row.get("流程顺序"),
            "granularity": row.get("模型粒度"),
            "internal_operations": row.get("内部操作说明"),
            "branch_role": row.get("主支线角色"),
            "source_refs": row.get("工艺骨架依据"),
            "include": True,
            "expand_mode": row.get("图中展开方式") or "as_modelled",
            "display_label": row.get("中文名称") or row.get("模型过程原名"),
            "en_name": row.get("模型过程原名"),
            "provenance": "模型过程:" + process_id,
        }
        for process_id, row in source["processes"].items()
    ]
    spec["edges"] = []
    for flow_id, flow in source["flows"].items():
        decision = source["drawing"].get(flow_id, {})
        if decision and decision.get("对象类型") != "Edge":
            raise ValueError(f"绘图确认 {flow_id}: 对象类型应为 Edge")
        show_model = flow.get("是否显示")
        show_drawing = decision.get("是否显示")
        if show_model in {"是", "否"} and show_drawing in {"是", "否"} and show_model != show_drawing:
            raise ValueError(f"流 {flow_id}: 模型表与绘图确认表的是否显示互相矛盾")
        kind = flow.get("值类型")
        if kind not in PHYSICAL | NONPHYSICAL:
            raise ValueError(f"流 {flow_id}: 值类型须明确为物理/环境交换或方法接口")
        include = kind in PHYSICAL and show_model != "否" and show_drawing != "否"
        if include and not numeric(flow.get("图示代表值")):
            raise ValueError(f"流 {flow_id}: 图示代表值缺少已计算的数值缓存")
        if include and (not flow.get("采用单位") or not flow.get("采用计量基准")):
            raise ValueError(f"流 {flow_id}: 单位或计量基准缺失")
        source_id, target_id = endpoints(flow, process_ids) if kind in PHYSICAL else (None, None)
        value = flow.get("图示代表值")
        spec["edges"].append({
            "edge_id": "edge_" + str(flow_id),
            "flow_id": flow_id,
            "name": flow.get("流名称"),
            "category": flow.get("流类别"),
            "value_type": kind,
            "source_node_id": source_id,
            "target_node_id": target_id,
            "source_value": value,
            "representative_value": abs(value) if numeric(value) else None,
            "display_value": None,
            "unit": flow.get("采用单位"),
            "basis": flow.get("采用计量基准"),
            "sign_convention": flow.get("符号约定"),
            "method_use_status": flow.get("方法采用状态"),
            "branch_role": decision.get("主支线ID"),
            "visual_role": kind if include else None,
            "route_class": "input_vertical" if source_id == "EXT_IN" else "output_vertical" if target_id == "EXT_OUT" else None,
            "visual_group_id": None,
            "include": include,
            "display_label": decision.get("显示标签") or flow.get("显示标签") or flow.get("流名称"),
            "provenance": "模型流与公式:" + str(flow_id),
        })
    return spec


def audit_spec(spec, workbook):
    issues = []
    freeze_errors, _ = validate_data(Path(workbook), freeze=True)
    issues.extend("数据功能冻结: " + item for item in freeze_errors)
    source = load_snapshot(workbook)
    if spec.get("metadata", {}).get("source_content_sha256") != source["sha256"]:
        issues.append("绘图规格引用的冻结 Excel 内容指纹不一致；请重新编译")
    p = source["project"]
    for key, label in (("functional_unit", "功能单位"), ("system_boundary", "系统边界"),
                       ("modeling_mode", "建模模式"), ("inventory_version", "冻结版本")):
        if spec.get("metadata", {}).get(key) != p.get(label):
            issues.append(f"metadata.{key} 与冻结 Excel 不一致")
    node_ids = {node.get("process_id") for node in spec.get("nodes", [])}
    if node_ids != set(source["processes"]):
        issues.append("模型过程覆盖与冻结 Excel 不一致")
    for node in spec.get("nodes", []):
        process = source["processes"].get(node.get("process_id"))
        if not process:
            continue
        for key, expected in (("workshop_id", process.get("功能车间")),
                              ("name", process.get("中文名称") or process.get("模型过程原名")),
                              ("sequence", process.get("流程顺序"))):
            if node.get(key) != expected:
                issues.append(f"过程 {node.get('process_id')}: {key} 与冻结 Excel 不一致")
    edges = spec.get("edges", [])
    flow_ids = [edge.get("flow_id") for edge in edges]
    if len(flow_ids) != len(set(flow_ids)) or set(flow_ids) != set(source["flows"]):
        issues.append("绘图规格流 ID 集合与冻结 Excel 不一致")
    for edge in edges:
        flow_id = edge.get("flow_id")
        flow = source["flows"].get(flow_id)
        if not flow:
            continue
        if flow.get("值类型") in PHYSICAL:
            try:
                source_id, target_id = endpoints(flow, set(source["processes"]))
            except ValueError as exc:
                issues.append(str(exc))
                continue
        else:
            source_id, target_id = None, None
        for key, expected in (("source_node_id", source_id), ("target_node_id", target_id),
                              ("unit", flow.get("采用单位")), ("basis", flow.get("采用计量基准")),
                              ("value_type", flow.get("值类型"))):
            if edge.get(key) != expected:
                issues.append(f"流 {flow_id}: {key} 与冻结 Excel 不一致")
        value = flow.get("图示代表值")
        expected_value = abs(value) if numeric(value) else None
        if edge.get("representative_value") != expected_value:
            issues.append(f"流 {flow_id}: 图示代表值与冻结 Excel 不一致")
        decision = source["drawing"].get(flow_id)
        expected_include = (flow.get("值类型") in PHYSICAL and flow.get("是否显示") != "否"
                            and (not decision or decision.get("是否显示") != "否"))
        if edge.get("include") != expected_include:
            issues.append(f"流 {flow_id}: include 与 Excel 的显示决定不一致")
        if edge.get("include") and flow.get("值类型") not in PHYSICAL:
            issues.append(f"流 {flow_id}: 方法负担或未知类型不得作为物理边显示")
        formula_row = source["drawing_formulas"].get(flow_id)
        if decision and decision.get("对象类型") == "Edge":
            for label, expected in (("来源过程ID", source_id), ("目标过程ID", target_id)):
                val = decision.get(label)
                refined = (isinstance(val, str) and expected not in ("EXT_IN", "EXT_OUT")
                           and val.startswith(str(expected) + "-OP"))
                if val not in (None, "") and val != expected and not refined:
                    issues.append(f"绘图确认 {flow_id}: {label} 与模型端点不一致")
            for label, expected in (("代表值", expected_value), ("单位", flow.get("采用单位"))):
                cached = decision.get(label)
                if cached not in (None, "", expected):
                    issues.append(f"绘图确认 {flow_id}: {label} 的公式缓存与模型不一致")
            for key in ("代表值", "单位"):
                formula = formula_row.get(key) if formula_row else None
                if not isinstance(formula, str) or not formula.startswith("=IF("):
                    issues.append(f"绘图确认 {flow_id}: {key} 必须为模型联动公式，不能手填")
    return issues


def main():
    parser = argparse.ArgumentParser(description="Compile a draft spec from a frozen LCI workbook")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    template = json.loads(args.template.read_text(encoding="utf-8"))
    try:
        spec = compile_spec(args.workbook, template)
        issues = audit_spec(spec, args.workbook)
        if issues:
            raise ValueError("；".join(issues))
    except ValueError as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(spec, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"DRAFT: {len(spec['nodes'])} processes, {len(spec['edges'])} flows; {args.output}")


if __name__ == "__main__":
    main()
